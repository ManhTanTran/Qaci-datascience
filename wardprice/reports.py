"""CSV, Excel, JSON, and dated daily-report exports."""

import csv
import html
import json
from datetime import datetime, time, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .database import (
    latest_valuation_rows,
    listing_rows,
    listing_statistics,
    stats,
    tracked_listing_count,
)
from .national_registry import national_coverage
from .price_warehouse import national_listing_rows, national_statistics_rows

COLUMNS = [
    "ward_code",
    "ward_name",
    "province",
    "property_type",
    "mean_vnd_m2",
    "min_vnd_m2",
    "max_vnd_m2",
    "data_type",
    "quality_gate",
    "confidence",
    "estimation_basis",
    "grounding_used",
    "source_urls",
    "collected_at",
    "model",
]
NUMERIC_COLUMNS = {"mean_vnd_m2", "min_vnd_m2", "max_vnd_m2"}
LISTING_COLUMNS = [
    "listing_id", "ward_code", "ward_name", "province", "property_type", "price_vnd",
    "area_m2", "price_per_m2", "address", "source_url", "source_domain", "listed_date",
    "collected_at", "last_seen_at", "evidence_status", "evidence_excerpt",
]
STAT_COLUMNS = [
    "ward_code", "ward_name", "province", "property_type", "listing_count",
    "mean_price_per_m2", "median_price_per_m2", "min_price_per_m2", "max_price_per_m2",
    "data_quality", "last_updated",
]
NATIONAL_LISTING_COLUMNS = [
    "listing_id", "administrative_code", "ward_name", "province", "target_rank",
    "title", "property_type", "price_vnd", "area_m2", "area_basis", "price_per_m2",
    "address", "source_website", "source_url", "source_domain", "published_at",
    "collected_at", "last_seen_at", "evidence_status", "validation_status",
    "included_in_stats",
]
NATIONAL_STAT_COLUMNS = [
    "administrative_code", "ward_name", "province", "target_rank", "property_type",
    "area_basis", "snapshot_date", "sample_count", "mean_price_per_m2",
    "median_price_per_m2", "min_price_per_m2", "max_price_per_m2",
    "calculation_window_days", "latest_observation_at", "source_quality",
    "calculation_version",
]


def protect_cell(obj):
    """Prefix spreadsheet formula-like text before writing CSV or Excel cells."""
    value = str(obj if obj is not None else "")
    if value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _daily_activity(db, local_now: datetime) -> dict:
    local_start = datetime.combine(local_now.date(), time.min, tzinfo=local_now.tzinfo)
    local_end = local_start + timedelta(days=1)
    start_utc = local_start.astimezone(timezone.utc).isoformat(timespec="seconds")
    end_utc = local_end.astimezone(timezone.utc).isoformat(timespec="seconds")
    row = db.execute(
        """
        SELECT COUNT(*) AS runs,
               COALESCE(SUM(attempted), 0) AS wards_attempted,
               COALESCE(SUM(accepted), 0) AS price_rows_accepted,
               COALESCE(SUM(no_data), 0) AS no_data,
               COALESCE(SUM(failed), 0) AS failed
        FROM run_log
        WHERE started_at >= ? AND started_at < ?
        """,
        (start_utc, end_utc),
    ).fetchone()
    activity = {key: int(value) for key, value in dict(row).items()}
    listing_row = db.execute(
        """SELECT COUNT(*) AS listing_search_calls,
                  COALESCE(SUM(accepted_count), 0) AS listings_added,
                  COALESCE(SUM(updated_count), 0) AS listings_updated,
                  COALESCE(SUM(rejected_count), 0) AS listing_candidates_rejected,
                  COALESCE(SUM(limit_skipped_count), 0) AS listing_limit_skipped,
                  COALESCE(SUM(search_queries), 0) AS grounding_search_queries
           FROM listing_search_runs WHERE started_at >= ? AND started_at < ?""",
        (start_utc, end_utc),
    ).fetchone()
    activity.update({key: int(value) for key, value in dict(listing_row).items()})
    return activity


def _records(db) -> list[dict]:
    records = []
    for row in latest_valuation_rows(db):
        record = dict(row)
        record["source_urls"] = json.loads(record.pop("source_urls_json"))
        records.append(record)
    return records


def _write_workbook(
    path: Path,
    records: list[dict],
    summary: dict,
    activity: dict,
    listings: list[dict],
    ward_stats: list[dict],
    national_listings: list[dict],
    national_stats: list[dict],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Latest prices"
    sheet.append(COLUMNS)
    for record in records:
        values = []
        for column in COLUMNS:
            value = record.get(column)
            if column == "source_urls":
                value = "; ".join(value or [])
            elif column == "grounding_used":
                value = "Yes" if value else "No"
            elif column not in NUMERIC_COLUMNS:
                value = protect_cell(value)
            values.append(value)
        sheet.append(values)

    listing_sheet = workbook.create_sheet("Listings")
    listing_sheet.append(LISTING_COLUMNS)
    for record in listings:
        listing_sheet.append(
            [
                record.get(column)
                if column in {"price_vnd", "area_m2", "price_per_m2"}
                else protect_cell(record.get(column))
                for column in LISTING_COLUMNS
            ]
        )

    stats_sheet = workbook.create_sheet("Ward type statistics")
    stats_sheet.append(STAT_COLUMNS)
    for record in ward_stats:
        stats_sheet.append(
            [
                record.get(column)
                if column
                in {
                    "listing_count",
                    "mean_price_per_m2",
                    "median_price_per_m2",
                    "min_price_per_m2",
                    "max_price_per_m2",
                }
                else protect_cell(record.get(column))
                for column in STAT_COLUMNS
            ]
        )

    national_listing_sheet = workbook.create_sheet("National listings")
    national_listing_sheet.append(NATIONAL_LISTING_COLUMNS)
    for record in national_listings:
        national_listing_sheet.append(
            [
                record.get(column)
                if column in {"price_vnd", "area_m2", "price_per_m2", "target_rank", "included_in_stats"}
                else protect_cell(record.get(column))
                for column in NATIONAL_LISTING_COLUMNS
            ]
        )

    national_stats_sheet = workbook.create_sheet("National price stats")
    national_stats_sheet.append(NATIONAL_STAT_COLUMNS)
    for record in national_stats:
        national_stats_sheet.append(
            [
                record.get(column)
                if column in {
                    "sample_count", "mean_price_per_m2", "median_price_per_m2",
                    "min_price_per_m2", "max_price_per_m2", "target_rank",
                    "calculation_window_days",
                }
                else protect_cell(record.get(column))
                for column in NATIONAL_STAT_COLUMNS
            ]
        )

    summary_sheet = workbook.create_sheet("Daily summary")
    summary_sheet.append(["Metric", "Value"])
    summary_sheet.append(["Report date (local)", datetime.now().astimezone().date().isoformat()])
    summary_sheet.append(["Daily activity", ""])
    for key, value in activity.items():
        summary_sheet.append([key, value])
    summary_sheet.append(["Cumulative coverage", ""])
    for key, value in summary.items():
        rendered = json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value
        summary_sheet.append([protect_cell(key), protect_cell(rendered) if isinstance(rendered, str) else rendered])

    header_fill = PatternFill("solid", fgColor="17365D")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True)
        for column_cells in worksheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 64)
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column in {5, 6, 7} and cell.value is not None:
                cell.number_format = "#,##0"
            elif isinstance(cell.value, str):
                cell.number_format = "@"
    for worksheet in (listing_sheet, stats_sheet, national_listing_sheet, national_stats_sheet):
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00" if isinstance(cell.value, float) else "#,##0"
                elif isinstance(cell.value, str):
                    cell.number_format = "@"
    workbook.save(path)


def export(db, destination: Path):
    destination.mkdir(parents=True, exist_ok=True)
    records = _records(db)
    legacy_summary = stats(db)
    listings = listing_rows(db)
    ward_stats = listing_statistics(db)
    nationwide_listings = national_listing_rows(db)
    nationwide_stats = national_statistics_rows(db)
    coverage = national_coverage(db)
    summary = {
        **legacy_summary,
        "tracked_listings": tracked_listing_count(db),
        "eligible_listings": len(listings),
        "ward_type_groups": len(ward_stats),
        "national_registry_status": coverage["registry_status"],
        "national_administrative_units": coverage["administrative_coverage"]["commune_count"],
        "national_verified_price_units": coverage["price_coverage"]["units_with_verified_price"],
        "national_listing_count": len(nationwide_listings),
        "national_verified_listing_count": sum(
            1 for row in nationwide_listings if row["included_in_stats"]
        ),
        "national_observation_count": db.execute(
            "SELECT COUNT(*) FROM listing_observations"
        ).fetchone()[0],
        "national_quarantine_count": db.execute(
            "SELECT COUNT(*) FROM listing_quarantine"
        ).fetchone()[0],
        "meanings": {
            "legacy_valuations": "AI estimate or unverified ward-level record; excluded from listing stats",
            "listing_stats": "Python summaries from distinct current listing records only",
        },
    }
    local_now = datetime.now().astimezone()
    report_day = local_now.date().isoformat()
    activity = _daily_activity(db, local_now)

    json_path = destination / "ward_prices_latest.json"
    csv_path = destination / "ward_prices_latest.csv"
    xlsx_path = destination / "ward_prices_latest.xlsx"
    summary_path = destination / "summary.json"
    html_path = destination / "summary.html"
    daily_json_path = destination / f"daily_report_{report_day}.json"
    daily_html_path = destination / f"daily_report_{report_day}.html"
    listings_csv_path = destination / "listings_latest.csv"
    listings_json_path = destination / "listings_latest.json"
    statistics_csv_path = destination / "ward_type_statistics.csv"
    national_listings_csv_path = destination / "national_listings_latest.csv"
    national_listings_json_path = destination / "national_listings_latest.json"
    national_statistics_csv_path = destination / "national_price_statistics.csv"

    json_path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    daily_payload = {
        "report_date_local": report_day,
        "generated_at": local_now.isoformat(timespec="seconds"),
        "daily_activity": activity,
        "cumulative_summary": summary,
        "ward_type_statistics": ward_stats,
    }
    listings_json_path.write_text(json.dumps(listings, ensure_ascii=False, indent=2), encoding="utf-8")
    national_listings_json_path.write_text(
        json.dumps(nationwide_listings, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    daily_json_path.write_text(
        json.dumps(daily_payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(COLUMNS)
        for record in records:
            row = []
            for column in COLUMNS:
                value = record.get(column)
                if column == "source_urls":
                    value = "; ".join(value or [])
                elif column == "grounding_used":
                    value = "Yes" if value else "No"
                row.append(protect_cell(value))
            writer.writerow(row)

    for csv_path_new, columns, rows in (
        (listings_csv_path, LISTING_COLUMNS, listings),
        (statistics_csv_path, STAT_COLUMNS, ward_stats),
        (national_listings_csv_path, NATIONAL_LISTING_COLUMNS, nationwide_listings),
        (national_statistics_csv_path, NATIONAL_STAT_COLUMNS, nationwide_stats),
    ):
        with csv_path_new.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(columns)
            for record in rows:
                writer.writerow(
                    [
                        record.get(column)
                        if column in {"price_vnd", "area_m2", "price_per_m2", "listing_count",
                                      "mean_price_per_m2", "median_price_per_m2", "min_price_per_m2",
                                      "max_price_per_m2", "sample_count", "target_rank",
                                      "calculation_window_days", "included_in_stats"}
                        else protect_cell(record.get(column))
                        for column in columns
                    ]
                )

    _write_workbook(
        xlsx_path, records, summary, activity, listings, ward_stats,
        nationwide_listings, nationwide_stats,
    )
    body = f"""<!doctype html><html lang="vi"><meta charset="utf-8">
    <title>WardPriceDaily — báo cáo {html.escape(report_day)}</title>
    <body style="font-family:Arial,sans-serif;max-width:760px;line-height:1.6">
    <h2>Báo cáo tin rao BĐS theo phường — {html.escape(report_day)}</h2>
    <p>Trong ngày: {activity["listing_search_calls"]} lượt tìm kiếm; thêm {activity["listings_added"]} tin,
    làm mới {activity["listings_updated"]} tin; loại {activity["listing_candidates_rejected"]} ứng viên;
    bỏ qua {activity["listing_limit_skipped"]} tin do chạm giới hạn.
    Đã dùng {activity["grounding_search_queries"]} truy vấn Search grounding theo metadata.</p>
    <p>Đang theo dõi {summary["tracked_listings"]} / 500 tin; {summary["listing_api_calls_this_month"]}
    lượt request Gemini tháng này. Request count không phải số search query hoặc chi phí.</p>
    <p>Kho toàn quốc: registry {html.escape(str(summary['national_registry_status']))},
    {summary['national_administrative_units']} đơn vị hành chính hiện hành, {summary['national_listing_count']}
    listing, {summary['national_observation_count']} observations; {summary['national_verified_price_units']}
    địa bàn có giá xác minh. Thống kê chỉ dùng SOURCE_PAGE_CONFIRMED đã được người vận hành rà soát.
    Tệp Excel giữ riêng loại BĐS và cơ sở diện tích.</p>
    <p><strong>Chất lượng nguồn:</strong> trạng thái do Gemini tự khai không đủ để đưa vào thống kê.
    Candidate mặc định là UNVERIFIED; cần review ngoài model, nguồn chi tiết và bằng chứng giá/diện tích/địa bàn.</p>
    <p><strong>Dữ liệu legacy:</strong> {legacy_summary["ai_estimate_rows"]} AI_ESTIMATE và
    {legacy_summary["market_reference_unverified_rows"]} valuation cũ được giữ riêng, không tính vào thống kê listing.</p>
    <p>Chi tiết tin rao và thống kê nằm trong tệp Excel/CSV đính kèm.</p>
    </body></html>"""
    html_path.write_text(body, encoding="utf-8")
    daily_html_path.write_text(body, encoding="utf-8")
    return {
        "csv": csv_path,
        "xlsx": xlsx_path,
        "json": json_path,
        "summary": summary_path,
        "html": html_path,
        "daily_json": daily_json_path,
        "daily_html": daily_html_path,
        "listings_csv": listings_csv_path,
        "listings_json": listings_json_path,
        "statistics_csv": statistics_csv_path,
        "national_listings_csv": national_listings_csv_path,
        "national_listings_json": national_listings_json_path,
        "national_statistics_csv": national_statistics_csv_path,
    }
