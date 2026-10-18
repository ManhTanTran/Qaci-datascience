import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from wardprice.config import Settings
from wardprice.database import (
    active_wards,
    connect,
    import_wards,
    listing_call_count,
    listing_statistics,
    record_listing,
    save_results,
    tracked_listing_count,
)
from wardprice.gemini import extract_listing_json, make_listing_prompt
from wardprice.listing_pipeline import preview_listing_selection, run_listing_cycle
from wardprice.mailer import dry_run_report
from wardprice.registry import Ward, load_ward_whitelist
from wardprice.reports import export
from wardprice.validate import normalize_listing, normalize_result

W1 = Ward("001", "Phường A", "Tỉnh B", "")
W2 = Ward("002", "Phường C", "Tỉnh B", "")


def settings(**overrides):
    base = Settings(
        "fake-model", "", False, 2, 500, 30, 30, 7, 3, False, "", "", "",
        max_total_listings=500,
        max_new_listings_per_ward=10,
    )
    return replace(base, **overrides)


def listing(
    ward=W1,
    *,
    url="https://example.vn/listing/1",
    price=1_000_000_000,
    area=50,
    kind="DAT_O",
    address="123 Đường Mẫu, Phường A",
    date="2026-10-01",
    evidence="SEARCH_SNIPPET_ONLY",
    page_read=False,
):
    return {
        "ward_code": ward.code,
        "ward_name": ward.name,
        "province": ward.province,
        "property_type": kind,
        "price_vnd": price,
        "area_m2": area,
        "address": address,
        "source_url": url,
        "listed_date": date,
        "evidence_status": evidence,
        "source_page_read": page_read,
        "price_evidence": f"Giá rao {price} đồng",
        "area_evidence": f"Diện tích {area} m2",
        "evidence_excerpt": f"Tin rao bán giá {price} đồng, diện tích {area} m2 tại {address}",
    }


class FakeResearcher:
    def __init__(self, responses=None):
        self.responses = responses or {}
        self.calls = []

    def research_listings(self, ward, tracked, *, max_new_listings):
        self.calls.append((ward.code, [row["listing_id"] for row in tracked], max_new_listings))
        value = self.responses.get(ward.code, {"listings": []})
        if isinstance(value, Exception):
            raise value
        return (
            value.get("payload", {"listings": []}),
            value.get("grounded_urls", []),
            value.get("usage", {"grounding_used": True, "search_queries": 1}),
        )


class FakeGemini429(Exception):
    code = 429


class ListingPipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.db_path = self.root / "state.sqlite3"
        self.db = connect(self.db_path)
        import_wards(self.db, [W1, W2])

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def _candidate(self, raw, ward=W1, grounded=None):
        url = raw["source_url"]
        return normalize_listing(raw, ward, grounded_urls=[url] if grounded is None else grounded)

    def test_registry_loads_operator_ward_whitelist_and_optional_district(self):
        path = self.root / "wards.csv"
        path.write_text(
            "ward_code,ward_name,province,district\n"
            "LOCAL-001,Phường Mẫu,Thành phố Mẫu,Quận Cũ\n",
            encoding="utf-8",
        )
        [ward] = load_ward_whitelist(path)
        self.assertEqual((ward.code, ward.name, ward.province, ward.district),
                         ("LOCAL-001", "Phường Mẫu", "Thành phố Mẫu", "Quận Cũ"))

    def test_listing_limit_is_global_500_and_refresh_continues_at_capacity(self):
        first = listing()
        candidate = self._candidate(first)
        self.assertEqual(record_listing(self.db, candidate, "seed", max_total=500), "INSERTED")
        for index in range(499):
            url = f"https://example.vn/seed/{index}"
            self.db.execute(
                """INSERT INTO listings(
                    listing_id,identity_key,identity_fingerprint,ward_code,ward_name,province,
                    property_type,price_vnd,area_m2,price_per_m2,address,source_url,source_domain,
                    listed_date,collected_at,last_seen_at,evidence_status,evidence_excerpt,raw_payload
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    f"seed-{index}", url, f"fingerprint-{index}", W1.code, W1.name, W1.province,
                    "DAT_O", 1_000_000_000, 50, 20_000_000, f"Địa chỉ {index} tại Phường A", url,
                    "example.vn", "2026-10-01", "2026-10-01T00:00:00+00:00",
                    "2026-10-01T00:00:00+00:00", "SEARCH_SNIPPET_ONLY", "evidence excerpt",
                    "{}",
                ),
            )
        self.db.commit()
        refreshed = listing(price=1_100_000_000)
        newcomer = listing(url="https://example.vn/listing/new", address="456 Đường Mới, Phường A")
        bot = FakeResearcher({
            W1.code: {
                "payload": {"listings": [refreshed, newcomer]},
                "grounded_urls": [refreshed["source_url"], newcomer["source_url"]],
            }
        })
        result = run_listing_cycle(
            self.db, settings(max_wards_per_run=1), bot,
            selected_ward_codes=[W1.code], logger=lambda _: None,
        )
        self.assertEqual(result["tracked_listings"], 500)
        self.assertEqual(result["updated"], 1)
        self.assertEqual(result["limit_skipped"], 1)
        self.assertEqual(tracked_listing_count(self.db), 500)
        self.assertEqual(
            self.db.execute("SELECT price_vnd FROM listings WHERE listing_id=?", (candidate["listing_id"],)).fetchone()[0],
            1_100_000_000,
        )

    def test_duplicate_url_refreshes_one_listing_and_keeps_history(self):
        first = listing()
        second = listing(price=1_200_000_000)
        bot = FakeResearcher({
            W1.code: {"payload": {"listings": [first]}, "grounded_urls": [first["source_url"]]}
        })
        run_listing_cycle(self.db, settings(max_wards_per_run=1), bot,
                          selected_ward_codes=[W1.code], logger=lambda _: None)
        bot.responses[W1.code] = {
            "payload": {"listings": [second]}, "grounded_urls": [second["source_url"]]
        }
        outcome = run_listing_cycle(self.db, settings(max_wards_per_run=1), bot,
                                    selected_ward_codes=[W1.code], logger=lambda _: None)
        self.assertEqual(outcome["inserted"], 0)
        self.assertEqual(outcome["updated"], 1)
        self.assertEqual(tracked_listing_count(self.db), 1)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM listing_history").fetchone()[0], 2)

    def test_cross_post_duplicate_by_exact_listing_fingerprint(self):
        first = listing()
        second = listing(url="https://another.vn/realty/abc")
        bot = FakeResearcher({
            W1.code: {"payload": {"listings": [first]}, "grounded_urls": [first["source_url"]]}
        })
        run_listing_cycle(self.db, settings(max_wards_per_run=1), bot,
                          selected_ward_codes=[W1.code], logger=lambda _: None)
        bot.responses[W1.code] = {
            "payload": {"listings": [second]}, "grounded_urls": [second["source_url"]]
        }
        run_listing_cycle(self.db, settings(max_wards_per_run=1), bot,
                          selected_ward_codes=[W1.code], logger=lambda _: None)
        self.assertEqual(tracked_listing_count(self.db), 1)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM listing_sources").fetchone()[0], 2)
        revised_cross_post = listing(url=second["source_url"], price=1_250_000_000)
        bot.responses[W1.code] = {
            "payload": {"listings": [revised_cross_post]},
            "grounded_urls": [revised_cross_post["source_url"]],
        }
        outcome = run_listing_cycle(self.db, settings(max_wards_per_run=1), bot,
                                    selected_ward_codes=[W1.code], logger=lambda _: None)
        self.assertEqual(outcome["inserted"], 0)
        self.assertEqual(outcome["updated"], 1)
        self.assertEqual(tracked_listing_count(self.db), 1)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM listing_history").fetchone()[0], 3)

    def test_whitelist_rejects_other_ward_and_filter_outside_whitelist(self):
        outside = listing(W2, url="https://example.vn/outside")
        bot = FakeResearcher({
            W1.code: {"payload": {"listings": [outside]}, "grounded_urls": [outside["source_url"]]}
        })
        result = run_listing_cycle(self.db, settings(max_wards_per_run=1), bot,
                                   selected_ward_codes=[W1.code], logger=lambda _: None)
        self.assertEqual(result["rejected"], 1)
        self.assertEqual(tracked_listing_count(self.db), 0)
        with self.assertRaisesRegex(ValueError, "outside the imported whitelist"):
            run_listing_cycle(self.db, settings(), bot, selected_ward_codes=["999"], logger=lambda _: None)

    def test_optional_district_context_survives_database_resume_and_prompt(self):
        ward = Ward("004", "Phường Mẫu", "Thành phố Mẫu", "top_50.txt row 3", "Quận Cũ")
        import_wards(self.db, [ward])
        self.db.close()
        self.db = connect(self.db_path)
        resumed = active_wards(self.db)[0]
        self.assertEqual(resumed.district, "Quận Cũ")
        self.assertIn(
            "Phường Mẫu, Quận Cũ, Thành phố Mẫu",
            make_listing_prompt(resumed, [], max_new_listings=1),
        )

    def test_empty_gemini_response_is_no_listings_not_a_failure(self):
        self.assertEqual(extract_listing_json(""), {"listings": []})
        bot = FakeResearcher({
            W1.code: {"payload": {"listings": []}, "grounded_urls": [],
                      "usage": {"grounding_used": False, "search_queries": 0}}
        })
        outcome = run_listing_cycle(self.db, settings(max_wards_per_run=1), bot,
                                    selected_ward_codes=[W1.code], logger=lambda _: None)
        self.assertEqual(outcome["failed"], 0)
        self.assertEqual(outcome["no_listings"], 0)  # empty response with no citations is logged as NO_GROUNDING
        self.assertEqual(tracked_listing_count(self.db), 0)
        status = self.db.execute("SELECT status FROM listing_search_runs").fetchone()[0]
        self.assertEqual(status, "NO_GROUNDING")

    def test_listing_monthly_request_cap(self):
        bot = FakeResearcher()
        capped = settings(max_wards_per_run=2, max_calls_per_month=1)
        first = run_listing_cycle(self.db, capped, bot, logger=lambda _: None)
        second = run_listing_cycle(self.db, capped, bot, logger=lambda _: None)
        self.assertEqual(first["selected"], 1)
        self.assertEqual(second["selected"], 0)
        self.assertEqual(len(bot.calls), 1)
        self.assertEqual(listing_call_count(self.db), 1)

    def test_rate_limit_stops_remaining_ward_requests(self):
        bot = FakeResearcher({W1.code: FakeGemini429("429 RESOURCE_EXHAUSTED quota exceeded")})
        events = []
        outcome = run_listing_cycle(
            self.db,
            settings(max_wards_per_run=2, max_calls_per_month=2),
            bot,
            selected_ward_codes=[W1.code, W2.code],
            logger=None,
            progress_callback=events.append,
        )
        self.assertEqual([call[0] for call in bot.calls], [W1.code])
        self.assertTrue(outcome["stopped_after_429"])
        self.assertEqual(outcome["not_attempted"], 1)
        self.assertEqual(outcome["api_calls_this_month"], 1)
        self.assertIn("not_attempted", [event["stage"] for event in events])
        self.assertEqual(self.db.execute("SELECT attempted FROM run_log").fetchone()[0], 1)

    def test_non_429_error_does_not_stop_other_ward_requests(self):
        bot = FakeResearcher({W1.code: RuntimeError("temporary service failure")})
        outcome = run_listing_cycle(
            self.db,
            settings(max_wards_per_run=2, max_calls_per_month=2),
            bot,
            selected_ward_codes=[W1.code, W2.code],
            logger=None,
        )
        self.assertEqual([call[0] for call in bot.calls], [W1.code, W2.code])
        self.assertFalse(outcome["stopped_after_429"])
        self.assertEqual(outcome["not_attempted"], 0)

    def test_evidence_requires_grounding_and_source_page_flag(self):
        raw = listing()
        snippet = self._candidate(raw)
        self.assertEqual(snippet["evidence_status"], "SEARCH_SNIPPET_ONLY")
        page = listing(evidence="SOURCE_PAGE_CONFIRMED", page_read=False)
        self.assertEqual(self._candidate(page)["evidence_status"], "SEARCH_SNIPPET_ONLY")
        confirmed = listing(evidence="SOURCE_PAGE_CONFIRMED", page_read=True)
        self.assertEqual(self._candidate(confirmed)["evidence_status"], "SOURCE_PAGE_CONFIRMED")
        with self.assertRaisesRegex(ValueError, "grounding sources"):
            self._candidate(raw, grounded=[])

    def test_price_per_m2_is_calculated_in_python(self):
        normalized = self._candidate(listing(price=1_000_000_000, area=50))
        self.assertEqual(normalized["price_per_m2"], 20_000_000)

    def test_statistics_are_grouped_by_ward_and_property_type(self):
        for raw in (
            listing(url="https://example.vn/a", price=1_000_000_000, area=50),
            listing(url="https://example.vn/b", price=2_000_000_000, area=50),
            listing(url="https://example.vn/c", price=900_000_000, area=30, kind="NHA_RIENG"),
        ):
            candidate = self._candidate(raw)
            record_listing(self.db, candidate, "stats-test", max_total=500)
        summary = listing_statistics(self.db)
        land = next(row for row in summary if row["property_type"] == "DAT_O")
        house = next(row for row in summary if row["property_type"] == "NHA_RIENG")
        self.assertEqual(land["listing_count"], 2)
        self.assertEqual(land["mean_price_per_m2"], 30_000_000)
        self.assertEqual(land["median_price_per_m2"], 30_000_000)
        self.assertEqual(house["listing_count"], 1)
        self.assertEqual(house["mean_price_per_m2"], 30_000_000)

    def test_legacy_ai_estimates_are_excluded_from_listing_statistics(self):
        legacy = {
            "ward_code": W1.code,
            "ward_name": W1.name,
            "province": W1.province,
            "property_type": "DAT_O",
            "average_price_per_m2": 25_000_000,
            "min_estimated_price_per_m2": 20_000_000,
            "max_estimated_price_per_m2": 30_000_000,
            "data_type": "AI_ESTIMATE",
            "estimation_basis": "Legacy estimated data kept for historical comparison only.",
            "source_urls": [],
        }
        normalized = normalize_result(legacy, W1, grounding_used=False, model="legacy-test")
        save_results(self.db, "legacy-run", [normalized])
        self.assertEqual(listing_statistics(self.db), [])
        self.assertEqual(tracked_listing_count(self.db), 0)

    def test_database_resume_preserves_whitelist_listing_and_history(self):
        raw = listing()
        record_listing(self.db, self._candidate(raw), "resume", max_total=500)
        self.db.close()
        self.db = connect(self.db_path)
        self.assertEqual({ward.code for ward in active_wards(self.db)}, {W1.code, W2.code})
        self.assertEqual(tracked_listing_count(self.db), 1)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM listing_history").fetchone()[0], 1)

    def test_daily_refresh_updates_listing_and_appends_observation(self):
        original = listing()
        bot = FakeResearcher({
            W1.code: {"payload": {"listings": [original]}, "grounded_urls": [original["source_url"]]}
        })
        run_listing_cycle(self.db, settings(max_wards_per_run=1), bot,
                          selected_ward_codes=[W1.code], logger=lambda _: None)
        revised = listing(price=1_250_000_000)
        bot.responses[W1.code] = {
            "payload": {"listings": [revised]}, "grounded_urls": [revised["source_url"]]
        }
        run_listing_cycle(self.db, settings(max_wards_per_run=1), bot,
                          selected_ward_codes=[W1.code], logger=lambda _: None)
        self.assertEqual(tracked_listing_count(self.db), 1)
        self.assertEqual(self.db.execute("SELECT price_vnd FROM listings").fetchone()[0], 1_250_000_000)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM listing_history").fetchone()[0], 2)
        self.assertEqual(listing_call_count(self.db), 2)

    def test_preview_selection_is_read_only_and_respects_whitelist_and_monthly_cap(self):
        result = preview_listing_selection(
            self.db,
            settings(max_wards_per_run=2, max_calls_per_month=1),
            selected_ward_codes=[W2.code],
        )
        self.assertEqual([ward.code for ward in result["selected"]], [W2.code])
        self.assertEqual(result["capacity"], 1)
        self.assertEqual(result["calls_left"], 1)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM run_log").fetchone()[0], 0)
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM listing_search_runs").fetchone()[0], 0)

    def test_pipeline_reports_ordered_stage_progress_for_live_ui(self):
        events = []
        bot = FakeResearcher()
        result = run_listing_cycle(
            self.db,
            settings(max_wards_per_run=1),
            bot,
            selected_ward_codes=[W1.code],
            logger=None,
            progress_callback=events.append,
        )
        self.assertEqual(result["selected"], 1)
        self.assertEqual(events[0]["stage"], "preflight")
        self.assertEqual(events[-1]["stage"], "run_complete")
        progress_values = [event["progress"] for event in events if event["progress"] is not None]
        self.assertEqual(progress_values, sorted(progress_values))
        self.assertIn("gemini", [event["stage"] for event in events])
        self.assertIn("database", [event["stage"] for event in events])

    def test_xlsx_csv_export_lists_and_statistics(self):
        raw = listing(address="=HYPERLINK('x') address, Phường A")
        record_listing(self.db, self._candidate(raw), "export", max_total=500)
        files = export(self.db, self.root / "outputs")
        self.assertTrue(files["listings_csv"].exists())
        self.assertTrue(files["statistics_csv"].exists())
        self.assertTrue(files["listings_json"].exists())
        workbook = load_workbook(files["xlsx"], read_only=True, data_only=True)
        self.assertEqual(workbook["Listings"].max_row, 2)
        self.assertEqual(workbook["Ward type statistics"].cell(2, 5).value, 1)
        self.assertTrue(workbook["Listings"].cell(2, 9).value.startswith("'=HYPERLINK"))
        workbook.close()
        self.assertEqual(json.loads(files["listings_json"].read_text(encoding="utf-8"))[0]["price_per_m2"], 20_000_000)

    def test_gmail_dry_run_builds_message_without_smtp(self):
        files = export(self.db, self.root / "out")
        with patch("wardprice.mailer.smtplib.SMTP_SSL") as smtp:
            preview = dry_run_report(settings(), files, "Test report")
            self.assertEqual(preview["to"], "not-configured@example.invalid")
            self.assertIn("listings_latest.csv", preview["attachments"])
            smtp.assert_not_called()


if __name__ == "__main__":
    unittest.main()
