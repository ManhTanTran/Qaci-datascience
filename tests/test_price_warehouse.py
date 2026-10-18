"""Offline listing warehouse tests with source-safe synthetic records."""

from __future__ import annotations

import csv
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

from wardprice.database import connect
from wardprice.national_registry import (
    import_admin_file,
    lookup_ward_price,
    national_collection_priorities,
)
from wardprice.price_warehouse import (
    import_listing_payload,
    load_listing_payload,
    prepare_listing_import,
)
from wardprice.reports import export

ADMIN_FIXTURE = Path(__file__).parent / "fixtures" / "nso_admin_fixture.xml"


def _db(path):
    db = connect(path)
    import_admin_file(
        db,
        ADMIN_FIXTURE,
        as_of_date="2025-07-01",
        is_current=True,
        source_reference="https://official.example/admin",
    )
    return db


def _listing(
    *,
    title="Nhà phố trung tâm",
    url="https://listings.example/nhapho/1",
    price=8_000_000_000,
    area=100,
    property_type="NHA_RIENG",
    area_basis="FLOOR_AREA",
    ward_code="00001",
    ward_name="Phường Trung Tâm",
    evidence_status="UNVERIFIED",
):
    return {
        "administrative_code": ward_code,
        "ward_name": ward_name,
        "province": "Tỉnh Demo",
        "title": title,
        "property_type": property_type,
        "price_vnd": price,
        "area_m2": area,
        "area_basis": area_basis,
        "address": f"12 Đường Demo, {ward_name}, Tỉnh Demo",
        "source_url": url,
        "evidence_status": evidence_status,
        "price_evidence": "Giá bán được nêu trên trang nguồn.",
        "area_evidence": "Diện tích được nêu trên trang nguồn.",
        "address_evidence": "Địa chỉ thuộc Phường Trung Tâm.",
    }


def _verify_kwargs():
    return {
        "operator_verified": True,
        "verified_by": "test reviewer",
        "verification_reference": "offline-fixture-review",
    }


def test_unverified_claim_is_saved_as_candidate_and_excluded_from_statistics(tmp_path):
    db = _db(tmp_path / "warehouse.sqlite3")
    try:
        raw = _listing(evidence_status="SOURCE_PAGE_CONFIRMED")
        result = import_listing_payload(
            db,
            {"listings": [raw]},
            apply=True,
            collected_at=datetime(2026, 10, 13, tzinfo=timezone.utc),
        )
        stored = db.execute("SELECT evidence_status,validation_status,included_in_stats FROM listings").fetchone()
        lookup = lookup_ward_price(db, "00001", property_type="NHA_RIENG")
        assert result["unverified_candidates"] == 1
        assert stored["evidence_status"] == "UNVERIFIED"
        assert stored["validation_status"] == "CANDIDATE"
        assert stored["included_in_stats"] == 0
        assert lookup["status"] == "UNVERIFIED_ONLY"
        assert lookup["statistics"] == []
        fallback = lookup_ward_price(db, "00002", refresh=True)
        assert fallback["status"] == "EXTERNAL_PROVIDER_NOT_CONFIGURED"
        assert fallback["provider_network_called"] is False
        plan = national_collection_priorities(db, limit=10)
        assert plan["tier_counts"]["P1"]["unit_count"] == 1
        assert plan["tier_counts"]["P1"]["active_listing_count"] == 1
        assert plan["tier_counts"]["P2"]["unit_count"] == 2
    finally:
        db.close()


def test_listing_daily_observations_are_idempotent_and_stats_are_separated(tmp_path):
    db = _db(tmp_path / "warehouse.sqlite3")
    day_one = datetime(2026, 10, 13, tzinfo=timezone.utc)
    records = [
        _listing(title="Nhà A", url="https://listings.example/house/a", price=8_000_000_000),
        _listing(title="Nhà B", url="https://listings.example/house/b", price=9_000_000_000),
        _listing(
            title="Đất C",
            url="https://listings.example/land/c",
            price=2_000_000_000,
            area=100,
            property_type="DAT_O",
            area_basis="LAND_AREA",
        ),
    ]
    try:
        first = import_listing_payload(
            db,
            {"listings": records},
            apply=True,
            run_id="first-day",
            collected_at=day_one,
            **_verify_kwargs(),
        )
        duplicate = import_listing_payload(
            db,
            {"listings": records},
            apply=True,
            run_id="same-day-repeat",
            collected_at=day_one,
            **_verify_kwargs(),
        )
        assert first["inserted"] == 3
        assert first["verified_candidates"] == 3
        assert duplicate["duplicates_existing"] == 3
        assert duplicate["observations_added"] == 0
        assert db.execute("SELECT COUNT(*) FROM listings").fetchone()[0] == 3
        assert db.execute("SELECT COUNT(*) FROM listing_observations").fetchone()[0] == 3

        house = lookup_ward_price(
            db, "00001", property_type="NHA_RIENG", area_basis="FLOOR_AREA"
        )
        land = lookup_ward_price(
            db, "00001", property_type="DAT_O", area_basis="LAND_AREA"
        )
        assert house["status"] == "FOUND"
        assert house["statistics"][0]["sample_count"] == 2
        assert house["statistics"][0]["mean_price_per_m2"] == 85_000_000
        assert house["statistics"][0]["median_price_per_m2"] == 85_000_000
        assert land["statistics"][0]["sample_count"] == 1
        assert land["statistics"][0]["mean_price_per_m2"] == 20_000_000

        next_day = datetime(2026, 10, 14, tzinfo=timezone.utc)
        refreshed = import_listing_payload(
            db,
            {"listings": records},
            apply=True,
            run_id="next-day-refresh",
            collected_at=next_day,
            **_verify_kwargs(),
        )
        assert refreshed["updated"] == 3
        assert refreshed["observations_added"] == 3
        assert db.execute("SELECT COUNT(*) FROM listing_observations").fetchone()[0] == 6
    finally:
        db.close()


def test_stale_statistics_are_reported_without_provider_network(tmp_path):
    db = _db(tmp_path / "stale.sqlite3")
    stale_day = (datetime.now(timezone.utc).astimezone().date() - timedelta(days=90)).isoformat()
    try:
        unit_id = db.execute(
            "SELECT id FROM administrative_units WHERE administrative_code='00001'"
        ).fetchone()[0]
        db.execute(
            """INSERT INTO ward_price_statistics(
               administrative_unit_id,property_type,area_basis,snapshot_date,sample_count,
               mean_price_per_m2,median_price_per_m2,min_price_per_m2,max_price_per_m2,
               calculation_window_days,latest_observation_at,source_quality,calculation_version
               ) VALUES(?,'NHA_RIENG','FLOOR_AREA',?,2,80000000,80000000,70000000,
               90000000,90,?,'OPERATOR_VERIFIED','test-v1')""",
            (unit_id, stale_day, f"{stale_day}T12:00:00+00:00"),
        )
        db.commit()
        result = lookup_ward_price(
            db, "00001", property_type="NHA_RIENG", area_basis="FLOOR_AREA"
        )
        assert result["status"] == "STALE_DATA"
        assert result["snapshot_date"] == stale_day
        assert result["statistics"][0]["sample_count"] == 2
    finally:
        db.close()


def test_source_url_dedup_and_wrong_ward_quarantine(tmp_path):
    db = _db(tmp_path / "dedup.sqlite3")
    try:
        row = _listing(url="https://listings.example/house/1?utm_source=search")
        duplicate = dict(row, source_url="https://www.listings.example/house/1")
        wrong_ward = _listing(
            title="Sai địa bàn",
            url="https://listings.example/house/wrong",
            ward_name="Xã Trùng",
        )
        result = import_listing_payload(
            db,
            {"listings": [row, duplicate, wrong_ward]},
            apply=True,
            run_id="dedup-run",
            collected_at=datetime(2026, 10, 13, tzinfo=timezone.utc),
        )
        assert result["duplicates_exact"] == 1
        assert result["candidate_count"] == 1
        assert result["quarantined"] == 1
        assert db.execute("SELECT COUNT(*) FROM listings").fetchone()[0] == 1
        assert db.execute("SELECT COUNT(*) FROM listing_quarantine").fetchone()[0] == 1
        assert db.execute("SELECT COUNT(*) FROM listing_sources").fetchone()[0] == 1
    finally:
        db.close()


def test_500_observation_daily_cap_and_xlsx_export(tmp_path):
    db = _db(tmp_path / "cap.sqlite3")
    try:
        rows = [
            _listing(
                title=f"Listing {index}",
                url=f"https://listings.example/house/{index}",
                price=8_000_000_000 + index,
            )
            for index in range(501)
        ]
        result = import_listing_payload(
            db,
            {"listings": rows},
            apply=True,
            run_id="daily-cap",
            collected_at=datetime(2026, 10, 13, tzinfo=timezone.utc),
        )
        assert result["inserted"] == 500
        assert result["limited"] == 1
        assert db.execute("SELECT COUNT(*) FROM listing_observations").fetchone()[0] == 500

        files = export(db, tmp_path / "outputs")
        workbook = load_workbook(files["xlsx"], read_only=True)
        assert "National listings" in workbook.sheetnames
        assert "National price stats" in workbook.sheetnames
        assert Path(files["national_listings_csv"]).is_file()
        assert Path(files["national_listings_json"]).is_file()
        assert Path(files["national_statistics_csv"]).is_file()
        assert workbook["National listings"].max_row == 501
        workbook.close()
    finally:
        db.close()


def test_json_csv_and_xlsx_inputs_load_offline(tmp_path):
    row = _listing()
    json_path = tmp_path / "listings.json"
    json_path.write_text(json.dumps({"listings": [row]}, ensure_ascii=False), encoding="utf-8")

    csv_path = tmp_path / "listings.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(row))
        writer.writeheader()
        writer.writerow(row)

    xlsx_path = tmp_path / "listings.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Listings"
    worksheet.append(list(row))
    worksheet.append(list(row.values()))
    workbook.save(xlsx_path)
    workbook.close()

    db = _db(tmp_path / "input-formats.sqlite3")
    try:
        for path in (json_path, csv_path, xlsx_path):
            payload = load_listing_payload(path)
            report = prepare_listing_import(db, payload)
            assert report["received"] == 1
            assert report["candidate_count"] == 1
            assert report["unverified_candidate_count"] == 1
    finally:
        db.close()
