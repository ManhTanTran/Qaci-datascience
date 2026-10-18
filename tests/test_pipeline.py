import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from wardprice.config import Settings
from wardprice.database import choose_wards, connect, month_call_count, stats
from wardprice.gemini import make_prompt
from wardprice.pipeline import run_cycle
from wardprice.registry import Ward, load_verified_wards
from wardprice.reports import export, protect_cell
from wardprice.validate import extract_json, normalize_result

W = Ward("00082", "Phường Ví Dụ", "Thành phố Ví Dụ", "https://danhmuchanhchinh.nso.gov.vn/")


def settings():
    return Settings("fake-model", "", False, 20, 100, 450, 30, 7, 3, False, "", "", "")


def price(ward=W, kind="DAT_O"):
    return {
        "ward_code": ward.code,
        "ward_name": ward.name,
        "province": ward.province,
        "property_type": kind,
        "average_price_per_m2": 30_000_000,
        "min_estimated_price_per_m2": 20_000_000,
        "max_estimated_price_per_m2": 40_000_000,
        "data_type": "MARKET_REFERENCE_UNVERIFIED",
        "estimation_basis": "Ước tính theo mặt bằng khu vực demo giả định; không phải giá thật.",
        "source_urls": [],
    }


class FakeResearcher:
    def __init__(self):
        self.calls = []

    def research(self, ward):
        self.calls.append(ward.code)
        return (
            {"results": [price(ward), price(ward, "CHUNG_CU")]},
            False,
            {"prompt_tokens": 120, "output_tokens": 75},
        )


class TestWardPrice(unittest.TestCase):
    def test_ward_average_gemini_prompt_is_retired(self):
        with self.assertRaisesRegex(RuntimeError, "listing-level research"):
            make_prompt(W)

    def test_registry_skips_unverified_and_duplicates(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "wards.csv"
            p.write_text(
                "ward_code,ward_name,province,admin_verified,admin_source\n"
                "00082,Phường A,Hà Nội,true,https://example.com/source\n"
                "00082,Phường A,Hà Nội,true,https://example.com/source\n"
                "00083,Phường B,Hà Nội,false,https://example.com/source\n",
                encoding="utf-8",
            )
            wards, warnings = load_verified_wards(p)
            self.assertEqual(len(wards), 1)
            self.assertEqual(len(warnings), 2)

    def test_estimate_cannot_sneak_in_as_verified(self):
        p = price()
        p["data_type"] = "MARKET_REFERENCE_UNVERIFIED"
        x = normalize_result(p, W, grounding_used=False, model="fake")
        self.assertEqual(x["data_type"], "AI_ESTIMATE")
        self.assertEqual(x["quality_gate"], "RESEARCH_ONLY")
        p["data_type"] = "MARKET_VERIFIED"
        with self.assertRaises(ValueError):
            normalize_result(p, W, grounding_used=True, model="fake")

    def test_code_and_range_check(self):
        p = price()
        p["ward_code"] = "WRONG"
        with self.assertRaisesRegex(ValueError, "ward_code"):
            normalize_result(p, W, grounding_used=False, model="fake")
        p = price()
        p["min_estimated_price_per_m2"] = 50_000_000
        with self.assertRaisesRegex(ValueError, "range"):
            normalize_result(p, W, grounding_used=False, model="fake")

    def test_json_parser_rejects_unstructured(self):
        self.assertEqual(extract_json('```json\n{"results": []}\n```')["results"], [])
        with self.assertRaises(ValueError):
            extract_json('{"wrong": []}')

    def test_deduplicated_coverage_persistence_and_reports(self):
        with tempfile.TemporaryDirectory() as d:
            dbpath = Path(d) / "state.db"
            db = connect(dbpath)
            bot = FakeResearcher()
            result = run_cycle(db, [W], settings(), bot, logger=lambda s: None)
            self.assertEqual(result["accepted_rows"], 2)
            self.assertEqual(stats(db)["wards_with_any_research_price"], 1)
            self.assertEqual(stats(db)["price_rows_latest"], 2)
            self.assertEqual(stats(db)["market_verified_rows"], 0)
            self.assertEqual(month_call_count(db), 1)
            self.assertEqual(
                choose_wards(
                    db, [W], limit=20, target=500, refresh_days=30, retry_days=7, max_attempts=3
                ),
                [],
            )
            files = export(db, Path(d) / "out")
            self.assertTrue(files["csv"].exists())
            self.assertTrue(files["json"].exists())
            self.assertTrue(files["xlsx"].exists())
            self.assertTrue(files["daily_html"].exists())
            self.assertTrue(files["daily_json"].exists())
            data = json.loads(files["json"].read_text(encoding="utf-8"))
            self.assertEqual(len(data), 2)
            workbook = load_workbook(files["xlsx"], read_only=True, data_only=True)
            self.assertEqual(workbook["Latest prices"]["E2"].value, 30_000_000)
            self.assertEqual(workbook["Latest prices"].max_row, 3)
            workbook.close()
            db.close()
            again = connect(dbpath)
            try:
                self.assertEqual(stats(again)["wards_with_any_research_price"], 1)
            finally:
                again.close()

    def test_rejects_duplicate_type_in_one_response(self):
        class Invalid:
            def research(self, ward):
                return {"results": [price(ward), price(ward)]}, False, {}

        with tempfile.TemporaryDirectory() as d:
            db = connect(Path(d) / "state.db")
            outcome = run_cycle(db, [W], settings(), Invalid(), logger=lambda s: None)
            self.assertEqual(outcome["failed"], 1)
            self.assertEqual(stats(db)["wards_with_any_research_price"], 0)
            db.close()

    def test_maximum_unique_wards_and_monthly_gate(self):
        with tempfile.TemporaryDirectory() as d:
            db = connect(Path(d) / "state.db")
            w1 = Ward("001", "Phường A", "Tỉnh A", "https://example.org/")
            w2 = Ward("002", "Phường B", "Tỉnh A", "https://example.org/")
            w3 = Ward("003", "Phường C", "Tỉnh A", "https://example.org/")
            capped = replace(
                settings(), max_total_wards=2, max_wards_per_run=3, max_calls_per_month=2
            )
            selected = choose_wards(
                db,
                [w1, w2, w3],
                limit=3,
                target=capped.max_total_wards,
                refresh_days=30,
                retry_days=7,
                max_attempts=3,
            )
            self.assertEqual(len(selected), 2)
            bot = FakeResearcher()
            result = run_cycle(db, [w1, w2, w3], capped, bot, logger=lambda s: None)
            self.assertEqual(result["successful_wards"], 2)
            self.assertEqual(stats(db)["wards_with_any_research_price"], 2)
            # 2 Gemini calls used this month, so a second run cannot call more.
            result2 = run_cycle(db, [w1, w2, w3], capped, bot, logger=lambda s: None)
            self.assertEqual(result2["selected"], 0)
            self.assertEqual(len(bot.calls), 2)
            db.close()

    def test_csv_formula_safety(self):
        self.assertEqual(protect_cell("=1+2"), "'=1+2")
        self.assertEqual(protect_cell("phuong"), "phuong")

    def test_xlsx_protects_formula_like_model_text(self):
        with tempfile.TemporaryDirectory() as d:
            db = connect(Path(d) / "state.db")
            result = normalize_result(price(), W, grounding_used=False, model="fake")
            result["estimation_basis"] = '=HYPERLINK("https://example.org","click")'
            from wardprice.database import save_results

            save_results(db, "test-run", [result])
            files = export(db, Path(d) / "out")
            workbook = load_workbook(files["xlsx"], read_only=True, data_only=True)
            self.assertTrue(workbook["Latest prices"]["K2"].value.startswith("'=HYPERLINK("))
            workbook.close()
            db.close()


if __name__ == "__main__":
    unittest.main()
