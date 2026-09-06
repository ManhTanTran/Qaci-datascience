"""Deterministic, privacy-safe demo data for the DC5 analysis UI."""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

from credit_scoring.dc5.clustering import CLUSTER_FEATURES
from credit_scoring.dc5.data import feature_sets

SAMPLE_COLUMNS = (
    "fid",
    "age_group",
    "gender",
    "city",
    "household_type",
    "income_group",
    "is_kyc",
    "service_count_group",
    "tenure_group",
    "recency_group",
    "telco_contract_age_group",
    "telco_internet_usage_group",
    "telco_internet_trend_group",
    "telco_payment_method",
    "telco_monetary_group",
    "utility_recency_group",
    "utility_frequency_group",
    "utility_amount_group",
    "retail_monetary_group",
    "retail_frequency_group",
    "healthcare_monetary_group",
    "healthcare_frequency_group",
    "vaccine_monetary_group",
)

AGE_LABELS = np.array(["18-24", "25-30", "31-40", "41-50", "51-60", "61-70", "Trên 70"])
TENURE_LABELS = np.array(
    ["Dưới 3 tháng", "3-6 tháng", "6-12 tháng", "1-2 năm", "2-3 năm", "3-5 năm", "Trên 5 năm"]
)
RECENCY_LABELS = np.array(
    ["Dưới 30 ngày", "1-3 tháng", "3-6 tháng", "6-12 tháng", "1-2 năm", "2-3 năm", "Trên 3 năm"]
)
ORDINAL_GROUPS = np.array(["Nhóm 1", "Nhóm 2", "Nhóm 3", "Nhóm 4"])


def _sigmoid(values: np.ndarray) -> np.ndarray:
    return 1.0 / (1.0 + np.exp(-values))


def _nullable(values: np.ndarray, missing_rate: float, rng: np.random.Generator) -> np.ndarray:
    output = values.astype(object)
    output[rng.random(len(output)) < missing_rate] = None
    return output


def build_demo_frame(n_rows: int = 100_000, seed: int = 20260805) -> pd.DataFrame:
    """Build a one-row-per-synthetic-customer frame accepted by the DC5 pipeline."""

    if n_rows < 10:
        raise ValueError("n_rows must be at least 10 so both target classes can be generated.")
    rng = np.random.default_rng(seed)
    engagement = rng.normal(0, 1, n_rows)
    purchasing_power = rng.normal(0, 1, n_rows)
    age_ord = rng.choice(7, n_rows, p=[0.07, 0.14, 0.25, 0.22, 0.16, 0.10, 0.06])
    tenure_ord = np.clip(
        np.rint(2.8 + 0.55 * engagement + 0.18 * age_ord + rng.normal(0, 1.15, n_rows)),
        0,
        6,
    ).astype("int8")
    recency_ord = np.clip(
        np.rint(3.0 - 0.9 * engagement + rng.normal(0, 1.2, n_rows)), 0, 6
    ).astype("int8")
    service_count = np.clip(
        np.rint(1.6 + 0.65 * engagement + 0.25 * purchasing_power + rng.normal(0, 0.8, n_rows)),
        0,
        4,
    ).astype("int8")
    active_domains = np.clip(service_count + rng.binomial(2, 0.35, n_rows), 1, 5).astype("int8")

    has_telco = rng.binomial(1, 0.91, n_rows).astype("int8")
    has_retail = rng.binomial(1, _sigmoid(-0.2 + 0.55 * purchasing_power), n_rows).astype("int8")
    has_pharmacy = rng.binomial(1, _sigmoid(-0.4 + 0.15 * age_ord), n_rows).astype("int8")
    has_app = rng.binomial(1, _sigmoid(0.45 + 0.65 * engagement), n_rows).astype("int8")
    has_loyalty = rng.binomial(1, _sigmoid(-0.1 + 0.75 * engagement), n_rows).astype("int8")

    high_value_probability = _sigmoid(
        -2.15
        + 0.50 * purchasing_power
        + 0.32 * engagement
        + 0.20 * service_count
        + 0.12 * tenure_ord
        - 0.10 * recency_ord
    )
    high_value = rng.binomial(1, high_value_probability, n_rows).astype("int8")
    monetary_non_high = rng.choice([1, 2, 3], n_rows, p=[0.35, 0.40, 0.25])
    telco_monetary_ord = np.where(high_value == 1, 4, monetary_non_high).astype("int8")
    telco_monetary_ord = np.where(has_telco == 1, telco_monetary_ord, np.nan)

    gender = rng.choice(["Nữ", "Nam"], n_rows, p=[0.49, 0.51])
    cities = rng.choice(
        ["Hà Nội", "TP. Hồ Chí Minh", "Đà Nẵng", "Hải Phòng", "Cần Thơ", "Khác"],
        n_rows,
        p=[0.25, 0.30, 0.09, 0.08, 0.06, 0.22],
    )
    household = rng.choice(
        ["Nhà thường", "Chung cư", "Nhà trọ", "Công ty", "Cơ quan nhà nước"],
        n_rows,
        p=[0.57, 0.20, 0.18, 0.03, 0.02],
    )
    income_ord = np.clip(
        np.rint(2.0 + 0.95 * purchasing_power + rng.normal(0, 0.75, n_rows)), 0, 4
    ).astype("int8")
    income_labels = np.array(
        ["Dưới 5trđ", "5-10trđ", "Trên 10trđ đến 18trđ", "18-30trđ", "Trên 30trđ"]
    )

    frame = pd.DataFrame(
        {
            "user_id": [f"DEMO-{index:07d}" for index in range(1, n_rows + 1)],
            "fid": np.arange(9_000_000_001, 9_000_000_001 + n_rows, dtype=np.int64),
            "age_group": AGE_LABELS[age_ord],
            "age_group_ord": age_ord,
            "gender": gender,
            "city": cities,
            "household_type": household,
            "income_group": income_labels[income_ord],
            "income_band_est_ord": income_ord,
            "is_kyc": rng.binomial(1, 0.78, n_rows).astype(bool),
            "service_count_group": np.array(
                ["1 service", "1 service", "2-3 services", "2-3 services", "4+ services"]
            )[service_count],
            "active_domain_count_ord": active_domains,
            "tenure_group": TENURE_LABELS[tenure_ord],
            "tenure_group_ord": tenure_ord,
            "recency_group": RECENCY_LABELS[recency_ord],
            "recency_group_ord": recency_ord,
            "has_telco": has_telco,
            "has_retail": has_retail,
            "has_pharmacy": has_pharmacy,
            "has_app": has_app,
            "has_loyalty": has_loyalty,
            "telco_monetary_group_ord": telco_monetary_ord,
        }
    )

    numeric_features = sorted(
        ({feature for schema in feature_sets().values() for feature in schema} | set(CLUSTER_FEATURES))
        - set(frame.columns)
        - {"retail_product_group"}
    )
    for offset, feature in enumerate(numeric_features):
        signal = 1.8 + 0.45 * engagement + 0.30 * purchasing_power + 0.05 * offset
        frame[feature] = np.clip(np.rint(signal + rng.normal(0, 1.0, n_rows)), 0, 8).astype(
            "int8"
        )
    frame["retail_product_group"] = rng.choice(
        ["MOBILE", "ACCESSORY", "LAPTOP", "APPLIANCE"], n_rows
    )

    frame["telco_contract_age_group"] = TENURE_LABELS[
        np.clip(tenure_ord + rng.integers(-1, 2, n_rows), 0, 6)
    ]
    frame["telco_internet_usage_group"] = ORDINAL_GROUPS[
        np.clip(frame["telco_internet_usage_group_ord"].to_numpy(), 0, 3)
    ]
    frame["telco_internet_trend_group"] = ORDINAL_GROUPS[
        np.clip(frame["telco_internet_trend_group_ord"].to_numpy(), 0, 3)
    ]
    frame["telco_payment_method"] = rng.choice(
        ["Chuyển khoản", "Ví điện tử", "Tiền mặt", "Thẻ"], n_rows
    )
    frame["telco_monetary_group"] = np.where(
        np.isnan(telco_monetary_ord), None, ORDINAL_GROUPS[np.nan_to_num(telco_monetary_ord - 1).astype(int)]
    )
    frame["utility_recency_group"] = _nullable(RECENCY_LABELS[recency_ord], 0.32, rng)
    frame["utility_frequency_group"] = _nullable(ORDINAL_GROUPS[rng.integers(0, 4, n_rows)], 0.32, rng)
    frame["utility_amount_group"] = _nullable(ORDINAL_GROUPS[rng.integers(0, 4, n_rows)], 0.32, rng)
    frame["retail_monetary_group"] = _nullable(ORDINAL_GROUPS[rng.integers(0, 4, n_rows)], 0.36, rng)
    frame["retail_frequency_group"] = _nullable(ORDINAL_GROUPS[rng.integers(0, 4, n_rows)], 0.36, rng)
    frame["healthcare_monetary_group"] = _nullable(ORDINAL_GROUPS[rng.integers(0, 4, n_rows)], 0.42, rng)
    frame["healthcare_frequency_group"] = _nullable(ORDINAL_GROUPS[rng.integers(0, 4, n_rows)], 0.42, rng)
    frame["vaccine_monetary_group"] = _nullable(ORDINAL_GROUPS[rng.integers(0, 4, n_rows)], 0.58, rng)
    return frame


def write_demo_parquet(frame: pd.DataFrame, output_path: str | Path) -> Path:
    """Write the complete pipeline-ready frame as Parquet."""

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_parquet(path, index=False)
    return path.resolve()


def write_demo_workbook(frame: pd.DataFrame, output_path: str | Path) -> Path:
    """Write the 23-column reference view as a formatted Excel workbook."""

    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    workbook._fonts[0] = Font(name="Arial", size=10)  # Set the workbook-wide default font.
    readme = workbook.create_sheet("README")
    notes = [
        ("DC5 synthetic demo dataset",),
        ("Mục đích", "Dùng để thử giao diện và pipeline; không dùng để ra quyết định thực tế."),
        ("Số dòng", len(frame)),
        ("Nguồn schema", "Workbook mẫu DC5xQACI- DMDL KH FPT_20260805 (2).xlsx"),
        ("Dữ liệu", "100% synthetic; không sao chép khách hàng, PII hoặc credential."),
        ("Seed", 20260805),
        ("File chạy UI", "data_extracted/model_df_extracted.parquet"),
    ]
    for row in notes:
        readme.append(row)
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 90

    data = workbook.create_sheet("03. VC_Mẫu dữ liệu")
    header = []
    for value in SAMPLE_COLUMNS:
        cell = WriteOnlyCell(data, value=value)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        header.append(cell)
    data.append(header)
    data.freeze_panes = "A2"
    data.auto_filter.ref = f"A1:W{len(frame) + 1}"
    for row in frame.loc[:, SAMPLE_COLUMNS].itertuples(index=False, name=None):
        data.append(row)
    workbook.save(path)
    return path.resolve()
